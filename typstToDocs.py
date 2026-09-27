import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_element(name):
    return OxmlElement(name)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for callout boxes (in dxa)"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = create_element('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = create_element(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color for a table cell"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = create_element('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_cell_border_left(cell, color_hex, size="36"):
    """Set a thick left border for callouts (like Typst quotes)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = create_element('w:tcBorders')

    left = create_element('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), size)
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), color_hex)
    tcBorders.append(left)

    # Clear other borders
    for b in ['top', 'bottom', 'right']:
        node = create_element(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)

    tcPr.append(tcBorders)

def build_beautiful_doc(filename, title, subtitle, content_structure):
    """
    Generates a Typst-like beautiful Word document.
    content_structure is a list of tuples: [('heading', 'Text'), ('paragraph', 'Text'), ('callout', 'Text')]
    """
    doc = Document()

    # 1. Page Setup (Clean 1-inch margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette Constants (Typst Minimalist)
    COLOR_PRIMARY = RGBColor(0x1A, 0x3A, 0x5C)    # Deep Navy
    COLOR_SECONDARY = RGBColor(0x4A, 0x55, 0x68)  # Slate Gray
    COLOR_TEXT = RGBColor(0x2D, 0x37, 0x48)       # Charcoal Text
    HEX_CALLOUT_BG = "F7FAFC"                     # Very Light Grey/Blue
    HEX_CALLOUT_BORDER = "3182CE"                 # Typst Blue Accent

    # 2. Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run(subtitle)
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(24)
    p_div_border = create_element('w:pBdr')
    bottom_border = create_element('w:bottom')
    bottom_border.set(qn('w:val'), 'single')
    bottom_border.set(qn('w:sz'), '6')
    bottom_border.set(qn('w:space'), '1')
    bottom_border.set(qn('w:color'), 'E2E8F0')
    p_div_border.append(bottom_border)
    p_div._p.get_or_add_pPr().append(p_div_border)

    # 3. Inject Content
    for item_type, text in content_structure:
        if item_type == 'heading':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = COLOR_PRIMARY

        elif item_type == 'paragraph':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.25  # Elegant line spacing
            run = p.add_run(text)
            run.font.name = 'Arial'
            run.font.size = Pt(10.5)
            run.font.color.rgb = COLOR_TEXT

        elif item_type == 'callout':
            # Create a 1x1 table to act as a Typst styled block quote box
            table = doc.add_table(rows=1, cols=1)
            table.autofit = False
            table.columns[0].width = Inches(6.5)

            cell = table.cell(0, 0)
            set_cell_shading(cell, HEX_CALLOUT_BG)
            set_cell_border_left(cell, HEX_CALLOUT_BORDER, size="24")
            set_cell_margins(cell, top=140, bottom=140, left=200, right=140)

            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            c_run = cp.add_run(text)
            c_run.font.name = 'Arial'
            c_run.font.size = Pt(10)
            c_run.font.italic = True
            c_run.font.color.rgb = COLOR_TEXT

            # Spacer after table
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(0)
            p_space.paragraph_format.space_after = Pt(8)

    # Save Document
    os.makedirs('generated', exist_ok=True)
    target_path = os.path.join('generated', filename)
    doc.save(target_path)
    print(f"Success! Saved to {target_path}")

# --- EXAMPLE USAGE ---
if __name__ == "__main__":
    data = [
        ('heading', '1. Introduction to Automated Typography'),
        ('paragraph', 'This document layout leverages precise paragraph styling principles inspired by modern typesetting systems like Typst. By establishing rigid font hierarchies and strict color rules, we can consistently generate polished business output.'),
        ('callout', 'Note: Generating documents via an internal engine ensures optimal performance without overhead, giving your AI agents predictable control over structures.'),
        ('heading', '2. Key Design Standards'),
        ('paragraph', 'Every paragraph features structural line-spacing adjustments designed to maximize read-flow metrics. Say goodbye to the cluttered margins and bloated standard templates of the past.')
    ]

    build_beautiful_doc('beautiful_document.docx', 'Automated Modern Documents', 'A token-efficient framework for AI agents', data)
